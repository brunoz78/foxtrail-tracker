#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Verwaltungs-CLI fuer den Foxtrail-Tracker.

  manage.py init-db                       Schema anlegen
  manage.py seed [datei.json]             Erstbefuellung aus data/trails_seed.json (nur neue Trails)
  manage.py sync                          Abgleich mit foxtrail.ch (fuer Cron/systemd-Timer)
  manage.py zeitplan                      woechentlicher Abgleich als Dauerprozess (Docker, ohne systemd)
  manage.py create-user NAME [--admin]    Benutzer anlegen (Passwort wird abgefragt oder aus
                                          $FOXTRAIL_PASSWORD gelesen)
  manage.py set-password NAME             Passwort neu setzen
  manage.py list-users
  manage.py import-excel DATEI.xlsx       "Gemacht?"-Spalte aus der Excel-Uebersicht uebernehmen
  manage.py import-bestellungen DATEI     Bestellungen von foxtrail.ch (Kontoseite als HTML oder
                                          Text, "-" = stdin) als gemacht eintragen; ohne
                                          --schreiben nur Probelauf
  manage.py export-json [datei]           Komplette Liste inkl. eigener Eintraege als JSON sichern
  manage.py stats

Die Datenbank wird ueber $FOXTRAIL_DB gewaehlt (siehe foxtrail/config.py).
"""

import argparse
import getpass
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from foxtrail import bestellungen, config, db, sync, trails, users, zeitplan  # noqa: E402


def _pw(prompt="Passwort: "):
    pw = os.environ.get("FOXTRAIL_PASSWORD")
    if pw:
        return pw
    pw = getpass.getpass(prompt)
    if getpass.getpass("Wiederholen: ") != pw:
        sys.exit("Passwoerter stimmen nicht ueberein.")
    return pw


def cmd_init_db(a):
    with db.session() as conn:
        db.init_db(conn)
    print("DB bereit:", config.db_path())


def cmd_seed(a):
    with db.session() as conn:
        n = trails.seed_from_file(conn, a.datei)
    print(f"{n} Trails neu eingefuegt.")


def cmd_sync(a):
    with db.session() as conn:
        res = sync.run(conn, ausloeser=a.ausloeser)
    if res["ok"]:
        print(f"ok: {res['gefunden']} gefunden, {res['neu']} neu, {res['aktualisiert']} aktualisiert, "
              f"{res['reaktiviert']} reaktiviert, {res['archiviert']} archiviert, "
              f"{res['nicht_mehr_im_angebot']} gemachte nicht mehr im Angebot")
    else:
        print("FEHLER:", res["meldung"], file=sys.stderr)
        sys.exit(1)


def cmd_zeitplan(a):
    zeitplan.laufen(config.db_path())


def cmd_create_user(a):
    with db.session() as conn:
        try:
            rolle = "admin" if a.admin else ("lesen" if a.nur_lesen else "bearbeiten")
            users.add(conn, a.name, _pw(), rolle=rolle)
        except users.UserError as ex:
            sys.exit(str(ex))
    print(f"Benutzer {a.name} ({users.ROLLEN[rolle]}) angelegt.")


def cmd_set_password(a):
    with db.session() as conn:
        try:
            users.set_password(conn, a.name, _pw("Neues Passwort: "))
        except users.UserError as ex:
            sys.exit(str(ex))
    print("Passwort gesetzt.")


def cmd_list_users(a):
    with db.session() as conn:
        for u in users.list_users(conn):
            print(f"{u['username']:20} {users.ROLLEN[users.rolle(u)]:14} "
                  f"{'aktiv' if u['active'] else 'inaktiv':8} letzte Anmeldung: {u['last_login'] or '-'}")


def cmd_stats(a):
    with db.session() as conn:
        s = trails.stats(conn)
    for k, v in s.items():
        print(f"{k:12} {v}")


def cmd_export_json(a):
    with db.session() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM trails ORDER BY ort, name")]
    out = json.dumps({"exportiert": db.now(), "trails": rows}, ensure_ascii=False, indent=1)
    if a.datei:
        open(a.datei, "w", encoding="utf-8").write(out)
        print(f"{len(rows)} Trails nach {a.datei} exportiert.")
    else:
        print(out)


def cmd_import_bestellungen(a):
    """Kontoseite foxtrail.ch -> Deine Bestellungen: als "Webseite, vollstaendig" gespeichert
    oder den Seitentext in eine Datei kopiert. Zuordnung ueber den exakten Trail-Namen."""
    if a.datei == "-":
        inhalt = sys.stdin.read()
    else:
        with open(a.datei, encoding="utf-8", errors="replace") as fh:
            inhalt = fh.read()
    eintraege = bestellungen.parse(inhalt)
    if not eintraege:
        sys.exit("Keine Bestellungen gefunden - ist das die Seite 'Deine Bestellungen'?")
    with db.session() as conn:
        plan = bestellungen.zuordnen(conn, eintraege, benutzer=a.benutzer)
        breite = max(len(e["name"]) for e, *_ in plan) + 2
        for e, t, aktion, grund in plan:
            ziel = f"{t['ort']} | {t['name']}" if t else "-"
            pers = f"{e['personen']} Pers." if e["personen"] else "?"
            print(f"{e['datum'] or '??????????'}  {('Trail ' + e['name']).ljust(breite + 6)} {pers:>8}  "
                  f"-> {aktion:12s} {ziel}{('  (' + grund + ')') if grund else ''}")
        n_setzen = sum(1 for *_, a_, _ in plan if a_ == "setzen")
        n_erg = sum(1 for *_, a_, _ in plan if a_ == "ergaenzen")
        if not a.schreiben:
            print(f"\nProbelauf: {n_setzen} Trail(s) wuerden als gemacht eingetragen, {n_erg} ergaenzt. "
                  "Zum Schreiben --schreiben anhaengen.")
            conn.rollback()
            return
        res = bestellungen.anwenden(conn, plan, a.benutzer, None if a.ohne_fotos else config.foto_dir())
    print(f"\n{res['gesetzt']} Trail(s) als gemacht eingetragen, {res['ergaenzt']} ergaenzt, "
          f"{res['gekennzeichnet']} als Import gekennzeichnet, "
          f"{res['fotos']} Schlussfoto(s) geladen, {res['foto_fehler']} nicht ladbar.")


def cmd_import_excel(a):
    """Liest die mit dem Tracker verwandte Excel-Uebersicht (Spalten: Nr., Ort / Region,
    Trail-Name, ..., Gemacht?, Datum gemacht, Bemerkung) und uebernimmt die Eintraege.
    Zeilen unterhalb der Ueberschrift "Manuell ergaenzte Trails" werden als manuelle
    Trails angelegt."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl fehlt:  pip install openpyxl")
    wb = openpyxl.load_workbook(a.datei, data_only=True)
    ws = wb.active
    header_row = None
    for r in range(1, 20):
        vals = [str(c.value or "").strip() for c in ws[r]]
        if "Trail-Name" in vals:
            header_row = r
            col = {v: i for i, v in enumerate(vals)}
            break
    if header_row is None:
        sys.exit("Kopfzeile mit 'Trail-Name' nicht gefunden.")

    def g(row, key):
        i = col.get(key)
        return row[i] if i is not None and i < len(row) else None

    import re

    def norm(s):
        return re.sub(r"[^a-z0-9]", "", str(s or "").lower())

    manual_mode, done, created, skipped = False, 0, 0, []
    with db.session() as conn:
        alle = [dict(r) for r in conn.execute("SELECT * FROM trails")]
        by_key = {(norm(t["ort"]), norm(t["name"])): t for t in alle}
        go_by_ort = {norm(t["ort"]): t for t in alle if t["typ"] == "go"}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            first = str(row[0] or "")
            if first.startswith("Manuell erg"):
                manual_mode = True
                continue
            ort, name = str(g(row, "Ort / Region") or "").strip(), str(g(row, "Trail-Name") or "").strip()
            if not ort or not name:
                continue
            gemacht = str(g(row, "Gemacht?") or "").strip().lower() == "ja"
            datum = g(row, "Datum gemacht")
            datum = datum.strftime("%Y-%m-%d") if hasattr(datum, "strftime") else (str(datum or "").strip() or None)
            bem = str(g(row, "Bemerkung") or "").strip()
            form = {"gemacht": "1" if gemacht else "", "gemacht_datum": datum or "", "bemerkung": bem}
            mit = g(row, "Mitspieler")
            if mit:
                form["mitspieler"] = str(mit).split(".")[0]
            is_go = "GO" in str(g(row, "Typ") or "") or norm(name) == "digitaleschnitzeljagd"
            t = by_key.get((norm(ort), norm(name)))
            if t is None and is_go:
                t = go_by_ort.get(norm(ort))
            if t:
                if gemacht or bem:
                    trails.update_done(conn, t["id"], form, a.benutzer)
                    done += 1
            elif manual_mode:
                typ = trails.typ_aus_name(name, go="GO" in str(g(row, "Typ") or ""))
                form.update({"ort": ort, "name": name, "typ": typ,
                             "route": str(g(row, "Route / Beschreibung") or ""),
                             "dauer": str(g(row, "Dauer") or "")})
                trails.add_manual(conn, form, a.benutzer)
                created += 1
            else:
                skipped.append(f"{ort} | {name}")
    print(f"{done} Trails uebernommen, {created} manuell angelegt.")
    if skipped:
        print("Nicht zugeordnet (nicht in der DB):", *skipped, sep="\n  ")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("init-db").set_defaults(fn=cmd_init_db)
    s = sub.add_parser("seed"); s.add_argument("datei", nargs="?"); s.set_defaults(fn=cmd_seed)
    s = sub.add_parser("sync"); s.add_argument("--ausloeser", default="cli"); s.set_defaults(fn=cmd_sync)
    sub.add_parser("zeitplan").set_defaults(fn=cmd_zeitplan)
    s = sub.add_parser("create-user"); s.add_argument("name")
    g = s.add_mutually_exclusive_group()
    g.add_argument("--admin", action="store_true"); g.add_argument("--nur-lesen", action="store_true")
    s.set_defaults(fn=cmd_create_user)
    s = sub.add_parser("set-password"); s.add_argument("name"); s.set_defaults(fn=cmd_set_password)
    sub.add_parser("list-users").set_defaults(fn=cmd_list_users)
    sub.add_parser("stats").set_defaults(fn=cmd_stats)
    s = sub.add_parser("export-json"); s.add_argument("datei", nargs="?"); s.set_defaults(fn=cmd_export_json)
    s = sub.add_parser("import-excel"); s.add_argument("datei"); s.add_argument("--benutzer", default="import")
    s.set_defaults(fn=cmd_import_excel)
    s = sub.add_parser("import-bestellungen"); s.add_argument("datei", help="HTML/Text oder - fuer stdin")
    s.add_argument("--schreiben", action="store_true", help="wirklich eintragen (sonst Probelauf)")
    s.add_argument("--ohne-fotos", action="store_true", help="Schlussfotos nicht herunterladen")
    s.add_argument("--benutzer", default="import"); s.set_defaults(fn=cmd_import_bestellungen)
    a = p.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
